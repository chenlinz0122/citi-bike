# -*- coding: utf-8 -*-
from docx import Document

path = r'D:\OneDrive\Desktop\项目提交材料\期末报告-纽约公共自行车数据分析.docx'
doc = Document(path)
ps = doc.paragraphs

# 机器学习相关段落完整输出
for idx in [48, 50, 71, 77, 79, 85, 88, 99, 100]:
    print('===== 段落 %d =====' % idx)
    print(ps[idx].text)
    print()

print('##################################################')
print('表格总览：')
for ti, tb in enumerate(doc.tables):
    print('--- 表 %d：%d 行 x %d 列 ---' % (ti, len(tb.rows), len(tb.columns)))
    for row in tb.rows:
        print(' | '.join(c.text.strip().replace('\n', '␤') for c in row.cells))
