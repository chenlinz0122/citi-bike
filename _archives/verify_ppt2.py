# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Emu

p = r'E:\citi-bike\_archives\pptgen\答辩PPT-纽约公共自行车数据分析_v2.pptx'
prs = Presentation(p)

# 1) 备注全文抽查
for idx in (1, 10, 12, 17):
    s = prs.slides[idx - 1]
    ns = s.notes_slide.notes_text_frame.text.strip()
    print('=== P%d 备注长度 %d 字 ===' % (idx, len(ns)))
    print(ns[:80].replace('\n', ' / '))
    print('...尾部:', ns[-40:].replace('\n', ' / '))

# 2) 形状越界检查（超过画布 13.33 x 7.5 英寸）
W, H = 13.33, 7.5
bad = 0
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        try:
            l, t, w, h = sh.left, sh.top, sh.width, sh.height
        except Exception:
            continue
        if l is None:
            continue
        r = (l + w) / 914400.0
        b = (t + h) / 914400.0
        l_in, t_in = l / 914400.0, t / 914400.0
        if l_in < -0.02 or t_in < -0.02 or r > W + 0.02 or b > H + 0.02:
            print('越界 P%d: %.2f,%.2f -> %.2f,%.2f' % (i, l_in, t_in, r, b))
            bad += 1
print('越界形状数:', bad)
