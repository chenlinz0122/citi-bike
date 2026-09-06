# -*- coding: utf-8 -*-
from docx import Document
import re

p = r'D:\OneDrive\Desktop\项目提交材料\期末报告-纽约公共自行车数据分析.docx'
doc = Document(p)
imgs = [r for r in doc.part.rels.values() if 'image' in r.reltype]
print('嵌入图片数:', len(imgs))
paras = [x.text for x in doc.paragraphs]
caps = [t.strip() for t in paras if re.search(r'图\s*\d+', t)]
print('图题数量:', len(caps))
for c in caps:
    print(' -', c[:55])
# 表格数
print('表格数:', len(doc.tables))
# 关键指标抽查
full = '\n'.join(paras)
for kw in ['13.2', 'K=4', '0.9465', '0.2031', '79.5', '299', '160', '497.5']:
    print(kw, '存在' if kw in full else '缺失')
