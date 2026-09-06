# -*- coding: utf-8 -*-
from pptx import Presentation

p = r'E:\citi-bike\_archives\pptgen\答辩PPT-纽约公共自行车数据分析_v2.pptx'
prs = Presentation(p)
print('总页数:', len(prs.slides))
tot_img = 0
for i, s in enumerate(prs.slides, 1):
    imgs = sum(1 for sh in s.shapes if sh.shape_type == 13)
    tot_img += imgs
    pg = ''
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            t = sh.text_frame.text.strip()
            if '/' in t and '17' in t and len(t) < 15:
                pg = t
                break
    notes = ''
    try:
        ns = s.notes_slide
        notes = ns.notes_text_frame.text.strip()
    except Exception:
        notes = '(无备注)'
    notes_head = notes.split('\n')[0][:30] if notes else '(无备注)'
    print('P%-2d imgs=%d %s | 备注: %s' % (i, imgs, pg, notes_head))
print('图片总数:', tot_img)
