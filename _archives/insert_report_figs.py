# -*- coding: utf-8 -*-
"""期末报告插入 6 张机器学习图（图7~图12），备份后修改。"""
import os
import shutil
import docx
from docx.shared import Cm, Pt
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = r"D:\OneDrive\Desktop\项目提交材料\期末报告-纽约公共自行车数据分析.docx"
BAK = r"D:\OneDrive\Desktop\项目提交材料\_edit\期末报告-纽约公共自行车数据分析_加图前备份.docx"
FIG = r"E:\citi-bike\output\figures"
OUT = r"D:\OneDrive\Desktop\项目提交材料\期末报告-纽约公共自行车数据分析.docx"

os.makedirs(os.path.dirname(BAK), exist_ok=True)
shutil.copy2(SRC, BAK)
print("已备份:", BAK)

doc = docx.Document(SRC)

def find_para(prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise RuntimeError(f"未找到锚点段落: {prefix}")

def target_after_para(p):
    """返回该段落后应紧贴插入的元素：若紧跟表格则插到表格之后，否则插到段落之后。"""
    nxt = p._element.getnext()
    if nxt is not None and nxt.tag == qn('w:tbl'):
        return nxt
    return p._element

def make_pic_para(img, caption, width_cm):
    pic_p = doc.add_paragraph()
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.paragraph_format.space_before = Pt(6)
    pic_p.paragraph_format.space_after = Pt(2)
    run = pic_p.add_run()
    run.add_picture(os.path.join(FIG, img), width=Cm(width_cm))
    cap_p = doc.add_paragraph()
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_p.paragraph_format.space_after = Pt(10)
    r = cap_p.add_run(caption)
    r.font.size = Pt(10.5)
    return pic_p, cap_p

def insert_after(p, img, caption, width_cm):
    target = target_after_para(p)
    pic_p, cap_p = make_pic_para(img, caption, width_cm)
    target.addnext(cap_p._element)
    target.addnext(pic_p._element)
    print(f"  已插入: {caption}")

p_6_3_body = find_para("采用 KMeans 对各站点标准化特征聚类")
p_table3   = find_para("表3 站点聚类结果")
p_table4   = find_para("表4 站点客流与骑行时长预测模型对比")
p_6_4_body = find_para("客流预测最优为梯度提升")
p_baseline = find_para("朴素基线对比")

# 从后往前插入，避免锚点漂移
insert_after(p_baseline, "ml_fig6_duration_prediction_scatter.png",
             "图12 骑行时长预测最优模型真实值-预测值散点与残差分布", 15.0)
insert_after(p_6_4_body, "ml_fig9_demand_prediction_timeseries.png",
             "图11 客流预测最优模型在测试集上的真实/预测时序对比", 15.0)
insert_after(p_table4, "ml_fig7_duration_model_comparison.png",
             "图10 骑行时长预测五模型对比（R²、MAE、RMSE）", 15.0)
insert_after(p_table4, "ml_fig10_demand_model_comparison.png",
             "图9 站点小时客流预测五模型对比（R²、MAE、RMSE）", 15.0)
insert_after(p_table3, "ml_fig3_kmeans_radar_chart.png",
             "图8 四类站点聚类画像雷达图（标准化特征）", 12.0)
insert_after(p_6_3_body, "ml_fig1_kmeans_elbow_method.png",
             "图7 聚类数K选择的四项指标（肘部法则、轮廓系数、CH指数、DB指数）", 14.5)

doc.save(OUT)
print(f"\n完成！6 张图已插入 -> {OUT}")
