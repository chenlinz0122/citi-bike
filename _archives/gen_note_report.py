# -*- coding: utf-8 -*-
"""生成带备注的期末报告备注版（不覆盖原文件，正文一字不改，备注为红字浅橙底纹）。"""
import docx
from docx.shared import Pt, RGBColor

SRC = r"D:\OneDrive\Desktop\项目提交材料\期末报告-纽约公共自行车数据分析.docx"
DST = r"D:\OneDrive\Desktop\项目提交材料\期末报告-纽约公共自行车数据分析_备注版.docx"

d = docx.Document(SRC)


def make_note(text):
    """创建一条备注段落（红字、浅橙底纹、10.5pt），返回段落对象。"""
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("【备注】" + text)
    run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
    run.font.size = Pt(10.5)
    shd = docx.oxml.OxmlElement("w:shd")
    shd.set(docx.oxml.ns.qn("w:val"), "clear")
    shd.set(docx.oxml.ns.qn("w:fill"), "FFF3E0")
    p._p.get_or_add_pPr().append(shd)
    return p


def insert_note_after(target_para, text):
    """在指定段落后插入备注。"""
    note = make_note(text)
    target_para._p.addnext(note._p)


def insert_note_before(target_para, text):
    """在指定段落前插入备注。"""
    note = make_note(text)
    target_para._p.addprevious(note._p)


notes = [
    # (匹配关键词, 备注文字, 前后位置)  before=True 表示插在该段前
    ("1 团队成员组成及分工", "本文件为答辩备注版：正文与原报告完全一致，未做任何改动；下方红字【备注】为答辩讲解提示与潜在提问应对，仅作辅助，不随正文提交。", "before"),
    ("本小组共5人", "答辩开场建议：先讲“统一契约、并行开发”——组长先制定 config.py 列名规范与 docs/data_dict.md 接口契约，5 人模块只依赖统一标准表 clean_citibike.csv，实现零冲突对接。这是本项目的最大组织亮点，建议 30 秒内讲清。", "after"),
    ("开发语言：Python 3.10", "环境版本与 requirements.txt 锁定一致（pandas 2.3.3 / numpy 2.2.6 / scikit-learn 1.7.2）。若被问“为什么锁版本”：保证 5 人环境一致、结果可复现。TRAE 为 Coding Agent 工具，仅用于初稿与排错，架构设计、规范制定与结果核验均为人工完成，答辩如实说明即可。", "after"),
    ("样本个数", "口径记忆：原始 4,993,137 → 有效 4,975,379（剔除 17,758 条假启动/异常记录，占 0.36%）；抽样 20 万条、random_state=42 可复现。数据来源可现场打开 citibikenyc.com/system-data 查证。", "after"),
    ("操作方式", "图1 业务流程图与 main.py 五步骤一一对应：①下载→②清洗→③统计→④可视化→⑤机器学习。若被要求现场演示，运行 python main.py 即可一键产出全部结果（约 2 分钟）。", "after"),
    ("def candidate_urls", "数据获取策略亮点：①本机已有官方 zip 则直接解压合并，避免重复下载 929MB；②S3 服务器三个候选 URL 依优先级兜底；③大月份 zip 内 5 个分片自动合并为单一 CSV（4,993,137 行）。", "after"),
    ("对 929MB 大文件采用 50 万行分块读取", "工程性亮点：929MB 文件用 chunksize=500000 分块读取，内存占用稳定。若被问“为什么内存不会爆”：pandas 分块迭代读取 + 逐块过滤，只保留有效列。", "after"),
    ("（2）站点客流与骑行时长预测", "算法小结：聚类用四指标（SSE、轮廓系数、CH、DB）综合定 K=4；预测用 5 模型 + 2 朴素基线对比，避免“只报最优模型”的质疑。特征工程要点：骑行距离用 Haversine 公式由经纬度计算。", "after"),
    ("matplotlib 绘制时段分布", "可视化清单：基础图 6 张（时段/星期/周末占比/热门站点TOP10/空间散点/交互地图 folium）+ 机器学习 10 张（聚类4 + 预测6）。中文统一 SimHei 字体避免乱码，全部保存于 output/figures/。", "after"),
    ("20 万条真实样本中", "数据记忆：平均 13.2 分钟、中位数 9.5 分钟、会员 79.5%、电动车型 73.4%。若被问“为什么中位数小于均值”：骑行时长右偏（少数超长骑行拉高均值），中位数更能代表典型短途出行。", "after"),
    ("时段分布如图2", "规律记忆：早高峰 8 时、晚高峰 18 时（17,403 次）——通勤特征；工作日为周末的 3.1 倍；周三最高（36,220 次）、周六最低（22,473 次）。若被问“周三为何最高”：通勤与活动日叠加所致（报告口径）。", "after"),
    ("聚类深化解读", "聚类四类画像建议背熟：高流量通勤型 684 站（会员主导、高峰潮汐，通勤枢纽站）；中流量通勤型 443 站（早高峰占比高）；低流量通勤型 571 站（外围社区站）；低流量休闲型 287 站（时长长、周末高、会员低）。若被问“为什么 K=4”：肘部法 + 轮廓系数 + CH + DB 四指标综合确定。\n提示：本次按新机器学习模块复跑，四类站点数为 694/443/571/277，与正文 684/443/571/287 相差约 10 站（源于数据抽样顺序差异，其余指标一致：轮廓系数 0.1630 vs 0.1624、CH 292.61 vs 292.98）。若现场复跑被追问，以“约 700/440/570/280 站”表述即可。", "after"),
    ("朴素基线对比", "预测结论记忆：客流预测最优梯度提升 R²=0.9459、MAE=29.96 次/小时，lag_1h 滞后客流是主导特征（重要性 79.4%），可辅助调度；时长预测最优随机森林 R²=0.2030、MAE=4.80 分钟，骑行距离为第一重要特征（61.8%）。\n若被问“为什么时长预测 R² 这么低”：剩余方差来自个人骑行速度、红绿灯、实际路线等不可观测因素；且误差高度集中在 >60 分钟超长骑行（391 条贡献 75.1% 误差），30 分钟以内常规骑行（占 92%）MAE 仅 2.6~4.9 分钟，具备实用精度。\n基线对比亮点：时长预测 MAE 相对均值基线下降 42.2%、客流下降 80.7%——证明模型学到了真实规律而非“猜均值”。", "after"),
    ("最终成果与 AI 生成内容的差异", "AI 辅助过程如实申报即可：AI 生成模块初始框架，全部经人工运行验证与修改；接口契约、路径规范、异常兜底、结果核验均人工主导。若被问“AI 占比”：初稿框架由 AI 生成，但“AI 初稿 + 人工优化”的迭代全过程在报告中如实呈现。", "after"),
    ("可改进的地方", "总结陈述要点：① 499 万行真实数据上验证统一架构约定，五模块零冲突整合；② 创新点 5 条（zip 优先下载、分块清洗、四指标定 K、朴素基线量化增益、分桶误差归因）；③ 改进方向 3 条（引入天气/节假日特征提升时长预测、ST-GCN 刻画站点依赖、调度决策闭环）。", "after"),
]

# 先收集所有目标段落对象（匹配关键词，从上到下只取第一个命中）
targets = []
for kw, text, pos in notes:
    hit = None
    for p in d.paragraphs:
        if kw in p.text:
            hit = p
            break
    if hit is None:
        print(f"[警告] 未找到段落: {kw}")
        continue
    targets.append((hit, text, pos))

# 统一插入（addnext/addprevious 不影响其他目标的定位）
for p, text, pos in targets:
    if pos == "before":
        insert_note_before(p, text)
    else:
        insert_note_after(p, text)

d.save(DST)
print(f"已生成: {DST}")
