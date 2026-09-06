# -*- coding: utf-8 -*-
import os, zipfile, hashlib, glob
from docx import Document
from pptx import Presentation

desk = r'D:\OneDrive\Desktop\项目提交材料'
zhou = os.path.join(desk, '调研报告-周晨琳.docx')
final = os.path.join(desk, '期末报告-纽约公共自行车数据分析.docx')
ppt = os.path.join(desk, '答辩PPT-纽约公共自行车数据分析.pptx')
pkg = os.path.join(desk, '[01]周晨琳_陈嘉欣_张子千_童悦家_陈静萤_自行车.zip')

def sha(p):
    h = hashlib.sha256()
    with open(p, 'rb') as f:
        for chunk in iter(lambda: f.read(65536), b''):
            h.update(chunk)
    return h.hexdigest()

def docx_text(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for tb in doc.tables:
        for row in tb.rows:
            parts.append(' | '.join(c.text for c in row.cells))
    return '\n'.join(parts)

print('=' * 70)
print('【1】文件本体信息（时间戳/大小）')
for name, p in [('调研报告-周晨琳', zhou), ('期末报告', final), ('答辩PPT', ppt), ('提交包', pkg)]:
    st = os.stat(p)
    import time
    print('%-12s %8d 字节  修改时间 %s' % (name, st.st_size, time.strftime('%Y-%m-%d %H:%M:%S', time.localtime(st.st_mtime))))

print()
print('=' * 70)
print('【2】调研报告-周晨琳 内容核验')
zt = docx_text(zhou)
zhou_should = ['乔少杰', '郭洪飞', '柯日宏', '徐悦甡', '仝照民', '姜晓', '戢晓峰', '李福',
               'R²≈0.95', 'R²≈0.20', '1,985 个有效站点', '低流量休闲型', '0.36%', '工作日骑行量占比 75.8%']
zhou_not = ['SSE 14,833.6', 'SSE 13,754.9', '轮廓系数 0.160', '61.6%', '79.2%', '智研咨询', 'Citi Bike 官方系统数据',
            '766 站', 'MAE=33.4', 'MAE=4.6 分钟']
print('应存在（应为True）:')
for m in zhou_should:
    print('  ', m, '->', m in zt)
print('不应存在（应为True=已删除）:')
for m in zhou_not:
    print('  ', m, '->', m not in zt)

print()
print('=' * 70)
print('【3】期末报告 内容核验')
ft = docx_text(final)
f_should = ['图7', '图8', '图9', '图10', '图11', '图12',
            '13,754.9', '0.9465', '0.2031', '0.36%', '4,975,379', '4,993,137', 'Python 3.10']
f_not = ['Python 3.13', '14,833.6', '0.940']
print('应存在:')
for m in f_should:
    print('  ', m, '->', m in ft)
print('不应存在:')
for m in f_not:
    print('  ', m, '->', m not in ft)

print()
print('=' * 70)
print('【4】答辩PPT 内容核验')
prs = Presentation(ppt)
n_slides = len(prs.slides)
n_notes = 0
n_imgs = 0
ptxt = ''
for s in prs.slides:
    if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
        n_notes += 1
    for shp in s.shapes:
        if shp.shape_type == 13:
            n_imgs += 1
        if shp.has_text_frame:
            ptxt += shp.text_frame.text + '\n'
print('总页数:', n_slides, '| 含讲稿备注页:', n_notes, '| 嵌入图片数:', n_imgs)
p_should = ['K=4', '13,754.9', '0.9465', '0.2031', '0.36%', 'KMeans 站点聚类分析']
print('应存在:')
for m in p_should:
    print('  ', m, '->', m in ptxt)

print()
print('=' * 70)
print('【5】提交包 vs 桌面 一致性')
stage = r'E:\citi-bike\_archives\verify_final'
import shutil
if os.path.isdir(stage):
    shutil.rmtree(stage)
os.makedirs(stage)
with zipfile.ZipFile(pkg) as z:
    z.extractall(stage)
pairs = [
    (zhou, os.path.join(stage, '调研报告', '调研报告-周晨琳.docx'), '调研报告-周晨琳'),
    (os.path.join(desk, '调研报告-陈嘉欣.docx'), os.path.join(stage, '调研报告', '调研报告-陈嘉欣.docx'), '调研报告-陈嘉欣'),
    (final, os.path.join(stage, '小组期末报告.docx'), '期末报告'),
    (ppt, os.path.join(stage, '答辩PPT.pptx'), '答辩PPT'),
]
for a, b, name in pairs:
    if os.path.exists(a) and os.path.exists(b):
        ok = sha(a) == sha(b)
        print('  %-12s %s' % (name, '一致' if ok else '不一致!!'))
    else:
        print('  %-12s 缺失: %s / %s' % (name, os.path.exists(a), os.path.exists(b)))
shutil.rmtree(stage)

print()
print('=' * 70)
print('【6】全盘搜索其他 周晨琳 副本')
found = []
for root in [r'D:\OneDrive\Desktop', r'D:\OneDrive\Documents', r'C:\Users\Lenovo\Downloads', r'C:\Users\Lenovo\Documents', r'E:\citi-bike']:
    if not os.path.isdir(root):
        continue
    for dirpath, dirs, files in os.walk(root):
        # 跳过备份与解压缓存目录，避免噪音
        if '_edit' in dirpath or '_archives' in dirpath or 'pkg_stage' in dirpath or 'verify' in dirpath:
            continue
        for f in files:
            if '周晨琳' in f and f.lower().endswith('.docx'):
                found.append(os.path.join(dirpath, f))
if found:
    for f in found:
        st = os.stat(f)
        print('  %s (%d 字节)' % (f, st.st_size))
else:
    print('  除桌面主文件外无其他副本')
