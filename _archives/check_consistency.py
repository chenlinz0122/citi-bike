# -*- coding: utf-8 -*-
"""核对期末报告、PPT 中的数据与最新运行结果，并检查引用文件是否存在。"""
import os
import docx
from pptx import Presentation

BASE = r"E:\citi-bike"
DOC = r"D:\OneDrive\Desktop\项目提交材料\期末报告-纽约公共自行车数据分析.docx"
PPT = r"D:\OneDrive\Desktop\项目提交材料\答辩PPT-纽约公共自行车数据分析.pptx"

print("=" * 70)
print("一、期末报告关键数据核对（正文 vs 最新运行结果）")
print("=" * 70)

d = docx.Document(DOC)
full = "\n".join(p.text for p in d.paragraphs)

# 运行结果（2026-09-05 main.py 完整复跑）
run = {
    "平均时长13.2": "13.2 分钟" in full,
    "中位数9.5": "9.5 分钟" in full,
    "会员79.5%": "79.5%" in full,
    "游客20.5%": "20.5%" in full,
    "电动车型73.4%": "73.4%" in full,
    "晚高峰18时17403次": "17,403" in full,
    "工作日151663": "151,663" in full,
    "周末48337": "48,337" in full,
    "周末占比24.2%": "24.2%" in full,
    "周三36220": "36,220" in full,
    "周六最低22473": "22,473" in full,
    "Pier61起点752": "752" in full,
    "聚类K=4": "K=4" in full,
    "聚类684站": "684" in full,
    "聚类443站": "443" in full,
    "聚类287站": "287" in full,
    "聚类571站": "571" in full,
    "客流梯度提升0.9459": "0.9459" in full,
    "客流MAE29.96": "29.96" in full,
    "时长随机森林0.2030": "0.2030" in full,
    "时长MAE4.80": "4.80" in full,
    "距离特征61.8%": "61.8%" in full,
    "lag_1h主导79.4%": "79.4%" in full,
    "基线客流降80.7%": "80.7%" in full,
    "基线时长降42.2%": "42.2%" in full,
    "Python 3.10": "Python 3.10" in full,
    "499万原始": "4,993,137" in full,
    "497.5万有效": "4,975,379" in full,
}
print(f"{'检查项':<24} {'报告中':<6} 状态")
for k, ok in run.items():
    print(f"{k:<24} {'存在' if ok else '缺失':<6} {'✓' if ok else '✗ 需核对'}")

# 报告中引用的输出文件路径
print()
print("二、报告中引用的输出文件是否存在")
refs = [
    "output/统计结果.txt",
    "output/station_summary.csv",
    "output/figures/hour_distribution.png",
    "output/figures/weekday_distribution.png",
    "output/figures/weekend_pie.png",
    "output/figures/hot_station_bar.png",
    "output/figures/station_scatter.png",
    "output/figures/nyc_bike_interactive_map.html",
    "机器学习模块/output/model/ml_summary_report.txt",
    "机器学习模块/output/model/kmeans_cluster_insights.csv",
    "机器学习模块/output/model/kmeans_evaluation_metrics.csv",
    "机器学习模块/output/model/kmeans_cluster_summary.csv",
    "机器学习模块/output/model/duration_error_analysis.csv",
    "机器学习模块/output/model/demand_error_analysis.csv",
    "机器学习模块/output/model/duration_feature_importance.csv",
    "机器学习模块/output/model/demand_feature_importance.csv",
]
for r in refs:
    ok = os.path.exists(os.path.join(BASE, r))
    print(f"{'✓' if ok else '✗'} {r}")

print()
print("=" * 70)
print("三、PPT 关键数据核对（正文 vs 最新运行结果）")
print("=" * 70)
p = pptx = Presentation(PPT)
ptext = []
for slide in pptx.slides:
    for shape in slide.shapes:
        if shape.has_text_frame:
            ptext.append(shape.text_frame.text)
ppt_full = "\n".join(ptext)

ppt_checks = {
    "P10 K=4": "K=4" in ppt_full,
    "P10 1985站": "1985" in ppt_full,
    "P10 10维": "10维" in ppt_full or "10 维" in ppt_full,
    "P10 SSE 13,749.5": "13,749.5" in ppt_full,
    "P10 轮廓0.1624": "0.1624" in ppt_full,
    "P10 684站": "684站" in ppt_full,
    "P10 443站": "443站" in ppt_full,
    "P10 287站": "287站" in ppt_full,
    "P10 571站": "571站" in ppt_full,
    "P11 时长R² 0.2030": "0.2030" in ppt_full,
    "P11 客流R² 0.9459": "0.9459" in ppt_full,
    "P11 MAE 4.80/29.96": ("4.80" in ppt_full and "29.96" in ppt_full),
    "P11 距离61.8%": "61.8%" in ppt_full,
    "P11 lag_1h 79.4%": "79.4%" in ppt_full,
    "P11 基线降80.7%/42.2%": ("80.7%" in ppt_full and "42.2%" in ppt_full),
    "P13 客流0.9459": "0.9459" in ppt_full,
    "P14 客流0.9459": "0.9459" in ppt_full,
    "P4 Python 3.10": "Python 3.10" in ppt_full,
}
print(f"{'检查项':<28} {'PPT中':<6} 状态")
for k, ok in ppt_checks.items():
    print(f"{k:<28} {'存在' if ok else '缺失':<6} {'✓' if ok else '✗ 需核对'}")

# PPT 中的图片数量
print()
print("四、PPT 嵌入图片数量")
img_count = 0
for slide in pptx.slides:
    for shape in slide.shapes:
        if shape.shape_type == 13:  # PICTURE
            img_count += 1
print(f"PPT 共嵌入 {img_count} 张图片")
