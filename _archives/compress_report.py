# -*- coding: utf-8 -*-
import re
from docx import Document

path = r'D:\OneDrive\Desktop\项目提交材料\调研报告-周晨琳.docx'
doc = Document(path)
p = doc.paragraphs[18]
assert '③机器学习挖掘' in p.runs[0].text or 'R²≈0.95' in p.runs[0].text

new_text = ('结合调研，本项目将上述前沿技术路线落地于纽约公共自行车真实开放数据：'
            '获取约 499 万条行程记录，清洗剔除异常后保留 497.5 万条，并抽样 20 万条用于分析。'
            '统计发现，骑行以短途通勤为主——平均时长 13.2 分钟、中位数 9.5 分钟，会员占比 79.5%；'
            '出行呈早高峰 8 时、晚高峰 18 时的双峰形态，工作日骑行量约为周末的 3.1 倍。'
            '机器学习挖掘方面，KMeans 聚类将 1,985 个有效站点划分为高、中、低流量通勤型与低流量休闲型四类'
            '（分别约 694、443、571、277 站），前三类会员主导、通勤特征明显，休闲型站点周末出行占比偏高；'
            '客流预测以梯度提升最优（R²≈0.95、MAE≈29 次/小时），骑行时长预测以随机森林最优（R²≈0.20、MAE≈4.8 分钟），'
            '骑行距离与前 1 小时客流滞后项分别是最关键特征。'
            '最后用 matplotlib、Folium 输出时段/星期分布、热门站点 TOP10 与站点空间交互地图，'
            '实现“清洗—统计—建模—可视化”闭环，把前沿时空分析技术应用于真实城市开放数据。')

p.runs[0].text = new_text
doc.save(path)

# 统计压缩后全篇字数
doc2 = Document(path)
ps = [pp.text for pp in doc2.paragraphs if pp.text.strip()]
total = sum(len(re.sub(r'\s', '', t)) for t in ps)
print('新段落长度:', len(re.sub(r'\s', '', new_text)))
print('全篇字符数(去空白):', total)
