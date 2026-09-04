# -*- coding: utf-8 -*-
import docx

d = docx.Document(r'D:/OneDrive/Desktop/软件开发实践1-项目要求【2026.9】.docx')
with open(r'C:/Users/Lenovo/nyc_citibike/_req.txt', 'w', encoding='utf-8') as f:
    for p in d.paragraphs:
        if p.text.strip():
            f.write(p.text + '\n')
    for i, t in enumerate(d.tables):
        f.write(f'--- 表格{i} ---\n')
        for row in t.rows:
            f.write(' | '.join(c.text.strip() for c in row.cells) + '\n')
print('done')
