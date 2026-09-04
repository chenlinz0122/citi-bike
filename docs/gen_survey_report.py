# -*- coding: utf-8 -*-
"""从调研报告 md 生成 docx（宋体小四、行距20磅、标题黑体）。"""
import os, re, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
import config

MD = os.path.join(config.BASE_DIR, "docs", "调研报告-纽约公共自行车.md")
OUTS = [os.path.join(config.BASE_DIR, "docs", "调研报告-纽约公共自行车.docx"),
        r"D:\OneDrive\Desktop\项目提交材料\调研报告-纽约公共自行车.docx"]

doc = Document()

def add_run(p, text, cn="宋体", size=12, bold=False):
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), cn)
    r.font.size = Pt(size)
    r.bold = bold
    return r

def para(text, size=12, bold=False, cn="宋体", center=False, indent=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(20)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif indent:
        pf.first_line_indent = Pt(24)
    add_run(p, text, cn=cn, size=size, bold=bold)

lines = open(MD, encoding="utf-8").read().splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    i += 1
    if not line or line == "---":
        continue
    if line.startswith("# "):
        # 封面样式标题
        for _ in range(4):
            doc.add_paragraph()
        para("《软件开发实践1》调研报告", size=18, bold=True, cn="黑体", center=True)
        doc.add_paragraph()
        para("纽约公共自行车数据分析与可视化\n——前沿技术与产业应用调研",
             size=14, bold=True, cn="黑体", center=True)
        doc.add_paragraph()
        para("姓名：____________　　学号：____________", size=12, center=True, indent=False)
        para("杭州电子科技大学　　2026年9月", size=12, center=True, indent=False)
        doc.add_page_break()
    elif line.startswith("## "):
        title = line[3:].strip()
        if title == "调研报告":
            continue
        para(title, size=14, bold=True, cn="黑体", indent=False)
    else:
        text = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
        para(text)

for out in OUTS:
    try:
        doc.save(out)
        print("saved:", out)
    except PermissionError:
        print("被占用，跳过:", out)

# 字数统计（不含参考文献与空白）
body_text = "\n".join(l for l in lines if l.strip() and not l.startswith("[") and "citibikenyc" not in l)
body_text = body_text.split("## 参考文献")[0]
print("正文字数(含标点):", len(re.sub(r"\s", "", body_text)))
